"""
Service de parsing de CV : extrait le texte brut d'un fichier PDF ou DOCX.

Ce module ne fait QUE de l'extraction de texte. Le nettoyage, la
segmentation en sections et l'extraction d'entités (compétences,
expérience, etc.) seront gérés par preprocessor.py et extractor.py.
"""

from __future__ import annotations

import io
from enum import Enum

import fitz  # PyMuPDF
from docx import Document


class UnsupportedFileTypeError(Exception):
    """Levée quand le type de fichier n'est ni PDF ni DOCX."""


class EmptyDocumentError(Exception):
    """Levée quand aucun texte n'a pu être extrait du document."""


class SupportedFileType(str, Enum):
    PDF = "pdf"
    DOCX = "docx"


def _detect_file_type(filename: str) -> SupportedFileType:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return SupportedFileType.PDF
    if lower.endswith(".docx"):
        return SupportedFileType.DOCX
    raise UnsupportedFileTypeError(
        f"Type de fichier non supporté pour '{filename}'. Formats acceptés : .pdf, .docx"
    )


def _sort_blocks_by_reading_order(blocks: list[tuple], page_width: float) -> list[tuple]:
    """
    Trie des blocs de texte PyMuPDF selon un ordre de lecture qui gère
    le cas d'une mise en page à deux colonnes (fréquent sur les CV avec
    une bande latérale : contact, compétences, formation...).

    Chaque bloc est un tuple (x0, y0, x1, y1, text, block_no, block_type)
    tel que renvoyé par page.get_text("blocks").
    """
    # On ignore les blocs vides et les blocs d'image (block_type != 0)
    text_blocks = [b for b in blocks if b[6] == 0 and b[4].strip()]
    if len(text_blocks) < 2:
        return text_blocks

    # Centre horizontal de chaque bloc, trié de gauche à droite
    centers = sorted((b[0] + b[2]) / 2 for b in text_blocks)

    # On cherche le plus grand écart entre deux centres consécutifs :
    # c'est le candidat le plus probable pour une frontière de colonnes.
    gaps = [(centers[i + 1] - centers[i], centers[i]) for i in range(len(centers) - 1)]
    largest_gap, gap_position = max(gaps, key=lambda g: g[0])

    # Seuil : l'écart doit représenter au moins 10% de la largeur de page
    # pour être considéré comme une vraie séparation de colonnes, sinon
    # c'est juste un espacement normal entre deux blocs de la même colonne.
    MIN_GAP_RATIO = 0.10
    is_two_column_layout = largest_gap > page_width * MIN_GAP_RATIO

    if not is_two_column_layout:
        # Mise en page simple colonne : tri classique haut → bas
        return sorted(text_blocks, key=lambda b: (b[1], b[0]))

    split_x = gap_position + (largest_gap / 2)
    left_column = [b for b in text_blocks if (b[0] + b[2]) / 2 <= split_x]
    right_column = [b for b in text_blocks if (b[0] + b[2]) / 2 > split_x]

    # On exige au moins 2 blocs de chaque côté : sinon ce n'est probablement
    # pas une vraie mise en page à deux colonnes (juste un bloc isolé, un logo...)
    if len(left_column) < 2 or len(right_column) < 2:
        return sorted(text_blocks, key=lambda b: (b[1], b[0]))

    left_column.sort(key=lambda b: (b[1], b[0]))
    right_column.sort(key=lambda b: (b[1], b[0]))

    return left_column + right_column


def _extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extrait le texte de toutes les pages d'un PDF, en respectant l'ordre
    de lecture même sur une mise en page à deux colonnes.
    """
    text_parts: list[str] = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            blocks = page.get_text("blocks")
            ordered_blocks = _sort_blocks_by_reading_order(blocks, page.rect.width)
            page_text = "\n".join(b[4].strip() for b in ordered_blocks)
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_text_from_docx(file_bytes: bytes) -> str:
    """Extrait le texte des paragraphes ET des tableaux d'un DOCX."""
    document = Document(io.BytesIO(file_bytes))
    text_parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

    # Beaucoup de CV utilisent des tableaux pour la mise en page
    # (colonne compétences / colonne expérience) : on les parcourt aussi.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    return "\n".join(text_parts)


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Point d'entrée principal du service.

    Args:
        filename: nom original du fichier (utilisé pour détecter le type).
        file_bytes: contenu binaire du fichier.

    Returns:
        Le texte brut extrait du document.

    Raises:
        UnsupportedFileTypeError: si l'extension n'est pas .pdf ou .docx.
        EmptyDocumentError: si aucun texte n'a pu être extrait.
    """
    file_type = _detect_file_type(filename)

    if file_type == SupportedFileType.PDF:
        raw_text = _extract_text_from_pdf(file_bytes)
    else:
        raw_text = _extract_text_from_docx(file_bytes)

    cleaned = raw_text.strip()
    if not cleaned:
        raise EmptyDocumentError(
            "Aucun texte n'a pu être extrait du document. "
            "Il s'agit peut-être d'un scan (image) nécessitant de l'OCR."
        )

    return cleaned