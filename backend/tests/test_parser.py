import io

import fitz
import pytest
from docx import Document

from app.services.parser import (
    EmptyDocumentError,
    UnsupportedFileTypeError,
    _detect_file_type,
    _sort_blocks_by_reading_order,
    extract_text,
)


# --- Helpers pour construire des fichiers de test à la volée ---


def _make_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell_text in enumerate(row):
                table.cell(i, j).text = cell_text
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _make_two_column_pdf_bytes() -> bytes:
    """Construit un PDF minimal à deux colonnes, pour tester la
    détection de colonnes de bout en bout (pas seulement la fonction
    de tri isolée)."""
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    # Colonne gauche (ex: bande latérale contact/compétences)
    page.insert_text((50, 50), "Contact")
    page.insert_text((50, 80), "Info gauche")
    # Colonne droite (ex: contenu principal), largement séparée en x
    page.insert_text((350, 50), "Titre Principal")
    page.insert_text((350, 80), "Contenu droite")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


# --- Détection de type de fichier ---


def test_detect_file_type_pdf():
    assert _detect_file_type("cv.pdf").value == "pdf"


def test_detect_file_type_docx():
    assert _detect_file_type("cv.docx").value == "docx"


def test_unsupported_file_type_raises():
    with pytest.raises(UnsupportedFileTypeError):
        extract_text("cv.txt", b"contenu quelconque")


# --- Extraction DOCX ---


def test_extract_text_from_docx_paragraphs():
    content = _make_docx_bytes(["Jean Dupont", "Développeur Python", "5 ans d'expérience"])
    text = extract_text("cv_test.docx", content)
    assert "Jean Dupont" in text
    assert "Développeur Python" in text


def test_extract_text_from_docx_includes_tables():
    """Beaucoup de CV utilisent des tableaux pour la mise en page ;
    on vérifie que leur contenu n'est pas perdu."""
    content = _make_docx_bytes(
        paragraphs=["CV de Marie Curie"],
        table_rows=[["Compétence", "Niveau"], ["Python", "Expert"]],
    )
    text = extract_text("cv_tableau.docx", content)
    assert "Python" in text
    assert "Expert" in text


def test_empty_docx_raises():
    content = _make_docx_bytes([])
    with pytest.raises(EmptyDocumentError):
        extract_text("cv_vide.docx", content)


# --- Détection de colonnes (fonction isolée, données synthétiques) ---


def test_sort_blocks_detects_two_columns():
    blocks = [
        (10, 10, 100, 30, "Contact\n", 0, 0),
        (10, 40, 100, 60, "Info gauche\n", 1, 0),
        (300, 10, 400, 30, "Titre Principal\n", 2, 0),
        (300, 40, 400, 60, "Contenu droite\n", 3, 0),
    ]
    result = _sort_blocks_by_reading_order(blocks, page_width=600)
    texts = [b[4] for b in result]
    assert texts == ["Contact\n", "Info gauche\n", "Titre Principal\n", "Contenu droite\n"]


def test_sort_blocks_single_column_no_reorder():
    blocks = [
        (50, 10, 200, 30, "Ligne 1\n", 0, 0),
        (50, 40, 200, 60, "Ligne 2\n", 1, 0),
        (50, 70, 200, 90, "Ligne 3\n", 2, 0),
    ]
    result = _sort_blocks_by_reading_order(blocks, page_width=600)
    texts = [b[4] for b in result]
    assert texts == ["Ligne 1\n", "Ligne 2\n", "Ligne 3\n"]


def test_sort_blocks_ignores_isolated_block_as_false_column():
    """Un bloc isolé (ex: un logo) très éloigné en x ne doit pas être
    interprété comme une deuxième colonne à part entière (garde-fou :
    au moins 2 blocs de chaque côté)."""
    blocks = [
        (50, 10, 150, 30, "A\n", 0, 0),
        (50, 40, 150, 60, "B\n", 1, 0),
        (500, 10, 550, 30, "Logo\n", 2, 0),
    ]
    expected = sorted(blocks, key=lambda b: (b[1], b[0]))
    result = _sort_blocks_by_reading_order(blocks, page_width=600)
    assert result == expected


def test_sort_blocks_ignores_image_blocks():
    """block_type == 1 signifie image : à exclure du texte extrait."""
    blocks = [
        (10, 10, 100, 30, "Texte\n", 0, 0),
        (10, 40, 100, 60, "", 1, 1),  # bloc image
    ]
    result = _sort_blocks_by_reading_order(blocks, page_width=600)
    assert len(result) == 1
    assert result[0][4] == "Texte\n"


# --- Test d'intégration : extraction PDF réelle à deux colonnes ---


def test_extract_text_from_two_column_pdf_respects_reading_order():
    pdf_bytes = _make_two_column_pdf_bytes()
    text = extract_text("cv_deux_colonnes.pdf", pdf_bytes)
    # La colonne de gauche doit apparaître avant la colonne de droite
    assert text.index("Contact") < text.index("Titre Principal")
    assert text.index("Info gauche") < text.index("Contenu droite")